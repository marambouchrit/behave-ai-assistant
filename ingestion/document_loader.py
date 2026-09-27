"""
document_loader.py
==================
Chargement et extraction du texte depuis les guides utilisateurs BeHave.

Chaque fichier est découpé en *segments* structurels, qui portent la position
réelle du texte dans le document :
  - PDF  : un segment par page            → metadata["page"]    (1-based)
  - DOCX : un segment par section de titre → metadata["section"] (chemin des
           titres Heading N, ex. "Création de Clients > Génération de Template")
  - TXT  : un seul segment

Le chunker découpe ensuite chaque segment ; page et section sont ainsi
propagées à chaque chunk et servent de source affichée à l'utilisateur.

Le module BeHave n'est PAS déduit ici : il vient du manifeste (manifest.py).

Formats supportés :
  - DOCX : fichiers Word  (python-docx)
  - PDF  : fichiers PDF   (PyPDF2)
  - TXT  : texte brut
"""

import re
import unicodedata
import logging
from pathlib import Path
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

# ── Patterns compilés une seule fois au chargement du module ─────────────────

_RE_TABLE_SEPARATOR = re.compile(r"\|[\s\-|]+\|$")
_RE_DECORATION_LINE = re.compile(r"[-_=•·\s]+$")
_RE_PAGE_NUMBER     = re.compile(r"[Pp]age\s+\d+\s*(sur|of|/)\s*\d+\.?$")
_RE_SIRYOS_HEADER   = re.compile(
    r"(DOCUMENT DE TRAVAIL"
    r"|GUIDE D['’]UTILISATEUR"
    r"|Réf\s*:\s*DTR-[A-Z0-9\-/\s]+"
    r")$",
    re.IGNORECASE,
)
_RE_BOLD_MARKDOWN   = re.compile(r"\*{1,2}([^*]+)\*{1,2}")
_RE_MULTI_SPACES    = re.compile(r"[ \t]{2,}")
# Points de suite d'une table des matières : "Introduction ........ 6"
_RE_TOC_LEADER      = re.compile(r"(?:\.\s?){8,}")

# Une ligne présente sur au moins cette part des pages d'un PDF est un
# en-tête / pied de page répété (bruit pour l'embedding et le LLM).
_REPEATED_LINE_PAGE_RATIO = 0.5
_REPEATED_LINE_MIN_PAGES  = 3

# Styles de titre Word : "Heading 2" (anglais) ou "Titre 2" (français).
_RE_HEADING_STYLE   = re.compile(r"^(?:Heading|Titre)\s*(\d)$", re.IGNORECASE)
_TITLE_STYLE_NAMES  = frozenset({"title", "titre"})

SECTION_SEPARATOR   = " > "


@dataclass
class Document:
    """Unité de base du pipeline RAG : texte + métadonnées."""
    page_content: str
    metadata: dict = field(default_factory=dict)

    def __repr__(self) -> str:
        preview = self.page_content[:80].replace("\n", " ")
        source  = self.metadata.get("source", "?")
        return f"Document(source={source!r}, preview={preview!r}...)"

    def with_context(self) -> str:
        """
        Texte précédé de sa position dans le document ("titre — section"),
        pour les modèles qui évaluent un chunk hors de son contexte
        (embedding, reranking).
        """
        header = " — ".join(
            value for value in (
                self.metadata.get("title") or self.metadata.get("source"),
                self.metadata.get("section"),
            )
            if value
        )
        return f"{header}\n{self.page_content}" if header else self.page_content


def clean_text(text: str) -> str:
    """
    Nettoie le texte extrait d'un PDF ou DOCX en une seule passe sur les lignes.
    Aucun re.sub global après le join — évite une deuxième traversée du texte.
    """
    if not text:
        return ""

    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    cleaned_lines: list[str] = []

    for line in text.split("\n"):
        line = line.strip()

        if not line:
            continue
        if _RE_TABLE_SEPARATOR.match(line):
            continue
        if _RE_DECORATION_LINE.match(line) and len(line) > 3:
            continue
        if _RE_PAGE_NUMBER.match(line):
            continue
        if _RE_SIRYOS_HEADER.match(line):
            continue
        if _RE_TOC_LEADER.search(line):
            continue

        line = _RE_BOLD_MARKDOWN.sub(r"\1", line)
        line = _RE_MULTI_SPACES.sub(" ", line)
        cleaned_lines.append(line)

    return "\n".join(cleaned_lines).strip()


def _base_metadata(path: Path, file_type: str) -> dict:
    return {
        "source":    path.name,
        "file_path": str(path),
        "file_type": file_type,
    }


# ── DOCX ─────────────────────────────────────────────────────────────────────

def _heading_level(paragraph) -> int | None:
    match = _RE_HEADING_STYLE.match(paragraph.style.name or "")
    return int(match.group(1)) if match else None


def load_docx(file_path: str) -> list[Document]:
    """
    Charge un .docx en un segment par section.

    Le corps est parcouru dans l'ordre réel (paragraphes et tableaux mêlés).
    Chaque titre "Heading N" ouvre une nouvelle section ; le paragraphe de
    style "Title", s'il existe, devient le titre du document.
    """
    import docx as python_docx
    from docx.table import Table

    path = Path(file_path)
    doc  = python_docx.Document(str(path))

    doc_title: str | None      = None
    heading_stack: list[tuple[int, str]] = []
    current_lines: list[str]   = []
    segments: list[Document]   = []

    def flush() -> None:
        content = clean_text("\n".join(current_lines))
        current_lines.clear()
        if not content:
            return
        metadata = _base_metadata(path, "docx")
        if heading_stack:
            metadata["section"] = SECTION_SEPARATOR.join(t for _, t in heading_stack)
        segments.append(Document(page_content=content, metadata=metadata))

    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    current_lines.append(" | ".join(cells))
            continue

        text = block.text.strip()
        if not text:
            continue

        if (block.style.name or "").lower() in _TITLE_STYLE_NAMES and doc_title is None:
            doc_title = text
            continue

        level = _heading_level(block)
        if level is not None:
            flush()
            heading_stack[:] = [(lvl, t) for lvl, t in heading_stack if lvl < level]
            heading_stack.append((level, text))

        current_lines.append(text)

    flush()

    if doc_title:
        for segment in segments:
            segment.metadata["title"] = doc_title

    if not segments:
        logger.warning("Fichier DOCX vide ou sans texte extractible : %s", path.name)
    return segments


# ── PDF ──────────────────────────────────────────────────────────────────────

def _normalize_line(line: str) -> str:
    return _RE_MULTI_SPACES.sub(" ", line).strip()


def _repeated_lines(pages_text: list[str]) -> set[str]:
    """
    Lignes présentes sur une large part des pages : en-têtes et pieds de page
    répétés. Détection statistique, indépendante du contenu du document.
    """
    if len(pages_text) < _REPEATED_LINE_MIN_PAGES:
        return set()

    page_count: dict[str, int] = {}
    for text in pages_text:
        for line in {_normalize_line(l) for l in text.split("\n")}:
            if line:
                page_count[line] = page_count.get(line, 0) + 1

    threshold = max(2, int(len(pages_text) * _REPEATED_LINE_PAGE_RATIO))
    return {line for line, count in page_count.items() if count >= threshold}


def load_pdf(file_path: str) -> list[Document]:
    """
    Charge un .pdf en un segment par page (numéro de page 1-based).
    Les en-têtes / pieds de page répétés sur la plupart des pages sont retirés.
    """
    import PyPDF2

    path     = Path(file_path)
    segments: list[Document] = []
    pages_text: list[tuple[int, str]] = []

    with open(str(path), "rb") as f:
        reader = PyPDF2.PdfReader(f)

        for page_index, page in enumerate(reader.pages):
            try:
                pages_text.append((page_index + 1, page.extract_text() or ""))
            except Exception as exc:
                logger.warning("Page %d ignorée (%s) : %s", page_index + 1, path.name, exc)

    repeated = _repeated_lines([text for _, text in pages_text])
    if repeated:
        logger.debug("%s : %d ligne(s) d'en-tête/pied répétées retirées", path.name, len(repeated))

    for page_number, text in pages_text:
        kept    = [l for l in text.split("\n") if _normalize_line(l) not in repeated]
        content = clean_text("\n".join(kept))
        if content:
            metadata = _base_metadata(path, "pdf")
            metadata["page"] = page_number
            segments.append(Document(page_content=content, metadata=metadata))

    if not segments:
        logger.warning("Fichier PDF vide ou sans texte extractible : %s", path.name)
    return segments


# ── TXT ──────────────────────────────────────────────────────────────────────

def load_txt(file_path: str) -> list[Document]:
    """Charge un .txt en un seul segment."""
    path    = Path(file_path)
    content = clean_text(path.read_text(encoding="utf-8", errors="ignore"))

    if not content:
        logger.warning("Fichier TXT vide ou sans texte extractible : %s", path.name)
        return []
    return [Document(page_content=content, metadata=_base_metadata(path, "txt"))]


_LOADERS: dict = {
    ".docx": load_docx,
    ".pdf":  load_pdf,
    ".txt":  load_txt,
}

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(_LOADERS)


def load_document(file_path: str | Path) -> list[Document]:
    """
    Charge un document en segments, avec le loader adapté à son extension.

    Raises:
        ValueError : extension non supportée.
    """
    path   = Path(file_path)
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)

    if loader is None:
        supported = ", ".join(sorted(_LOADERS))
        raise ValueError(
            f"Format non supporté : '{suffix}'. Formats acceptés : {supported}"
        )

    return loader(str(path))


def list_document_files(directories: list[Path]) -> list[Path]:
    """Liste les fichiers supportés des répertoires donnés (ignorés s'ils n'existent pas)."""
    files: list[Path] = []
    for directory in directories:
        if not directory.is_dir():
            logger.warning("Répertoire de documents introuvable, ignoré : %s", directory)
            continue
        files.extend(
            f for f in sorted(directory.iterdir())
            if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
        )
    return files


if __name__ == "__main__":
    from config import get_settings

    logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")

    settings = get_settings()
    for file in list_document_files([settings.documents_dir, settings.uploaded_docs_dir]):
        segments = load_document(file)
        logger.info(
            "• %s | %d segments | %d chars",
            file.name, len(segments), sum(len(s.page_content) for s in segments),
        )
